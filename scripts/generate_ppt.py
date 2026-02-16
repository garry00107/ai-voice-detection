from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import re

def create_presentation(md_file='FINAL_PRESENTATION.md', output_file='Team_GauravSulsule_AI_Voice_Detection.pptx'):
    prs = Presentation()

    # Read markdown content
    with open(md_file, 'r') as f:
        content = f.read()

    # Split into slides
    slides_content = re.split(r'\n---\n', content)

    for i, slide_text in enumerate(slides_content):
        slide_text = slide_text.strip()
        if not slide_text:
            continue

        # Determine slide layout
        if i == 0:  # Title Slide
            slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            subtitle = slide.placeholders[1]
            
            # Simple parsing for title slide
            lines = slide_text.split('\n')
            title_text = ""
            subtitle_text = ""
            
            for line in lines:
                if line.startswith('# '): # Main title
                     # Skip meta headers 
                    continue
                elif line.startswith('**AI Voice Detection System**'):
                     title_text = "AI Voice Detection System"
                elif line.startswith('*Defending'):
                     subtitle_text += line.strip('*') + "\n"
                elif line.startswith('- **Team:**'):
                     subtitle_text += line + "\n"
                elif line.startswith('- **Event:**'):
                     subtitle_text += line + "\n"
                elif line.startswith('- **Live Demo:**'):
                     subtitle_text += line + "\n"
            
            title.text = title_text if title_text else "AI Voice Detection System"
            subtitle.text = subtitle_text
            
        else:  # Content Slides
            slide_layout = prs.slide_layouts[1] # Title and Content
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            body = slide.placeholders[1]
            text_frame = body.text_frame
            text_frame.clear()  # Clear default text

            lines = slide_text.split('\n')
            
            # Extract title
            for line in lines:
                if line.startswith('## '):
                    title.text = line.replace('## ', '').strip()
                    break
            
            # Add content
            for line in lines:
                line = line.strip()
                if line.startswith('## '):
                    continue
                
                # HEADERS (H3) -> Bold paragraph
                if line.startswith('### '):
                    p = text_frame.add_paragraph()
                    p.text = line.replace('### ', '')
                    p.font.bold = True
                    p.font.size = Pt(20)
                    p.level = 0
                    p.space_before = Pt(12)
                
                # BULLETS
                elif line.startswith('- ') or line.startswith('* '):
                    p = text_frame.add_paragraph()
                    p.text = line[2:]
                    p.level = 1
                    
                # BLOCKQUOTES
                elif line.startswith('> '):
                    p = text_frame.add_paragraph()
                    p.text = line[2:]
                    p.font.italic = True
                    p.level = 1
                
                # TABLES (Simple text representation)
                elif line.startswith('|'):
                    # Skip separator lines
                    if '---' in line:
                        continue
                    # Format as bullet for simplicity
                    clean_line = line.strip('|').replace('|', '  |  ')
                    p = text_frame.add_paragraph()
                    p.text = clean_line
                    p.level = 1
                    p.font.name = 'Courier New'
                    p.font.size = Pt(12)

                # NORMAL TEXT
                elif line and not line.startswith('---'):
                    p = text_frame.add_paragraph()
                    p.text = line
                    p.level = 0

    prs.save(output_file)
    print(f"Presentation saved to {output_file}")

if __name__ == "__main__":
    create_presentation()
